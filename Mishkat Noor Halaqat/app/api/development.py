"""Scoped workflows for the nine additions in version 15."""
import io, json, secrets, zipfile
from datetime import date, datetime, timedelta
from typing import Literal
from fastapi import APIRouter, Depends, HTTPException, Request, UploadFile, File, Form
from fastapi.responses import Response
from pydantic import Field, model_validator, ValidationError
from sqlalchemy import select, func, delete, or_
from app.db.session import get_db
from app.models import (User, Mosque, AuthSession, AuditLog, Halaqa, HalaqaStudent, HalaqaProgress, HalaqaAccountRequest, StudentExam,
 QuranRecord, RingCalendar, CalendarException, AbsenceExcuse, StudentMessage, ThreadRead, Announcement, OrganizedExam, OrganizedResult, ImportPreview)
from app.api.halaqat import Input, StudentRequestInput, _require, _actor, _ring_ids_for, _user_pack, PORTAL_ROLES, STUDENT_ROLE, SUPERVISOR_ROLE, TEACHER_ROLE
from app.api.student_file import access_student, today
from app.api.routes import audit, audit_change
from app.services.calendar import meeting_day
from app.services.quran_catalog import SURAH_NAMES, QURAN_STATES

router=APIRouter(prefix='/api/halaqat')
STAFF={SUPERVISOR_ROLE,TEACHER_ROLE}

def ring_access(db,request,mid,rid,roles=STAFF):
    u=_require(request,db,mid,roles)
    if rid not in _ring_ids_for(u,mid,db): raise HTTPException(404,'الحلقة خارج صلاحياتك')
    return u

def active_students(db,u,mid):
    q=select(HalaqaStudent).where(HalaqaStudent.halaqa_id.in_(_ring_ids_for(u,mid,db)),HalaqaStudent.active.is_(True))
    if u.role==STUDENT_ROLE:q=q.where(HalaqaStudent.user_id==u.id)
    return db.scalars(q).all()

class MosqueInput(Input):
    name: str=Field(min_length=2,max_length=180)

@router.post('/organizations')
def create_mosque(data:MosqueInput,request:Request,db=Depends(get_db)):
    u=_require(request,db,1,{'owner'})
    if db.scalar(select(Mosque.id).where(Mosque.name==data.name)):raise HTTPException(409,'اسم المسجد موجود')
    row=Mosque(id=(db.scalar(select(func.max(Mosque.id))) or 0)+1,name=data.name);db.add(row);db.flush()
    audit(db,u.id,row.id,'mosque_created',data.name);db.commit()
    return {'id':row.id,'name':row.name}

@router.put('/mosques/{mid}/organization')
def rename_mosque(mid:int,data:MosqueInput,request:Request,db=Depends(get_db)):
    u=_require(request,db,mid,{'owner'});row=db.get(Mosque,mid)
    if db.scalar(select(Mosque.id).where(Mosque.name==data.name,Mosque.id!=mid)):raise HTTPException(409,'اسم المسجد موجود')
    audit_change(db,u.id,mid,'mosque_renamed',mid,{'name':row.name},data.model_dump());row.name=data.name;db.commit()
    return {'ok':True}

@router.get('/organizations/overview')
def mosque_overview(request:Request,db=Depends(get_db)):
    _require(request,db,1,{'owner'});out=[]
    for m in db.scalars(select(Mosque).order_by(Mosque.id)):
        ids=list(db.scalars(select(Halaqa.id).where(Halaqa.mosque_id==m.id,Halaqa.active.is_(True))))
        out.append({'id':m.id,'name':m.name,'rings':len(ids),'students':db.scalar(select(func.count(HalaqaStudent.id)).where(HalaqaStudent.halaqa_id.in_(ids),HalaqaStudent.active.is_(True))), 'pending':db.scalar(select(func.count(HalaqaAccountRequest.id)).where(HalaqaAccountRequest.mosque_id==m.id,HalaqaAccountRequest.status=='pending'))})
    return {'mosques':out}

class QuranInput(Input):
    status: Literal['unstarted','in_progress','memorized','consolidation','revision']
    notes: str=Field('',max_length=1000)

@router.put('/mosques/{mid}/students/{sid}/quran/{surah}')
def quran_save(mid:int,sid:int,surah:int,data:QuranInput,request:Request,db=Depends(get_db)):
    u,s=access_student(db,request,mid,sid,True)
    if not 1<=surah<=114:raise HTTPException(422,'رقم السورة غير صالح')
    row=db.scalar(select(QuranRecord).where(QuranRecord.student_id==sid,QuranRecord.surah==surah))
    before={'status':row.status,'notes':row.notes} if row else {'status':'unstarted','notes':''}
    if not row:row=QuranRecord(student_id=sid,surah=surah,approved_by=u.id);db.add(row)
    row.status=data.status;row.notes=data.notes;row.approved_by=u.id;row.updated_at=datetime.utcnow()
    audit_change(db,u.id,mid,'quran_updated',sid,before,data.model_dump()|{'surah':surah});db.commit()
    return {'ok':True}

class CalendarInput(Input):
    weekdays: list[int]=Field(default_factory=list,max_length=7)
    time_text: str=Field('',max_length=100)
    @model_validator(mode='after')
    def check(self):
        if len(set(self.weekdays))!=len(self.weekdays) or any(d not in range(7) for d in self.weekdays):raise ValueError('أيام الأسبوع غير صحيحة')
        return self

class ExceptionInput(Input):
    day: date
    held: bool=False
    reason: str=Field('',max_length=500)

@router.put('/mosques/{mid}/rings/{rid}/calendar')
def calendar_save(mid:int,rid:int,data:CalendarInput,request:Request,db=Depends(get_db)):
    u=ring_access(db,request,mid,rid,{SUPERVISOR_ROLE})
    row=db.get(RingCalendar,rid);before={'weekdays':json.loads(row.weekdays_json),'time_text':row.time_text} if row else {}
    if not row:row=RingCalendar(halaqa_id=rid);db.add(row)
    row.weekdays_json=json.dumps(data.weekdays);row.time_text=data.time_text
    audit_change(db,u.id,mid,'calendar_updated',rid,before,data.model_dump());db.commit();return {'ok':True}

@router.put('/mosques/{mid}/rings/{rid}/calendar-exception')
def calendar_exception(mid:int,rid:int,data:ExceptionInput,request:Request,db=Depends(get_db)):
    u=ring_access(db,request,mid,rid,{SUPERVISOR_ROLE})
    row=db.scalar(select(CalendarException).where(CalendarException.halaqa_id==rid,CalendarException.day==data.day))
    if not row:row=CalendarException(halaqa_id=rid,day=data.day);db.add(row)
    row.held=data.held;row.reason=data.reason
    audit(db,u.id,mid,'calendar_exception',json.dumps(data.model_dump(mode='json')|{'ring':rid},ensure_ascii=False));db.commit();return {'ok':True}

@router.delete('/mosques/{mid}/rings/{rid}/calendar-exception/{day}')
def calendar_exception_remove(mid:int,rid:int,day:date,request:Request,db=Depends(get_db)):
    u=ring_access(db,request,mid,rid,{SUPERVISOR_ROLE})
    db.execute(delete(CalendarException).where(CalendarException.halaqa_id==rid,CalendarException.day==day));audit(db,u.id,mid,'calendar_exception_removed',f'{rid}:{day}');db.commit();return {'ok':True}

class ExcuseInput(Input):
    day: date
    reason: str=Field(min_length=3,max_length=1500)

class ReviewInput(Input):
    status: Literal['approved','rejected']
    note: str=Field('',max_length=1000)

@router.post('/mosques/{mid}/students/{sid}/excuses')
def submit_excuse(mid:int,sid:int,data:ExcuseInput,request:Request,db=Depends(get_db)):
    u,s=access_student(db,request,mid,sid)
    if not s.active:raise HTTPException(409,'الطالب مؤرشف')
    if u.role!=STUDENT_ROLE:raise HTTPException(403,'يرفع العذر الطالب أو ولي أمره من واجهته')
    if not meeting_day(db,s.halaqa_id,data.day):raise HTTPException(422,'هذا اليوم ليس لقاءً في جدول الحلقة؛ راجع المشرف')
    row=db.scalar(select(AbsenceExcuse).where(AbsenceExcuse.student_id==sid,AbsenceExcuse.day==data.day))
    if row and row.status!='rejected':raise HTTPException(409,'يوجد عذر لهذا اليوم قيد المراجعة أو معتمد')
    if not row:row=AbsenceExcuse(student_id=sid,day=data.day,submitted_by=u.id);db.add(row)
    row.reason=data.reason;row.status='pending';row.submitted_by=u.id;row.review_note='';row.reviewed_by=None;row.updated_at=datetime.utcnow()
    audit(db,u.id,mid,'excuse_submitted',f'{sid}:{data.day}');db.commit();return {'ok':True}

@router.put('/mosques/{mid}/students/{sid}/excuses/{eid}')
def review_excuse(mid:int,sid:int,eid:int,data:ReviewInput,request:Request,db=Depends(get_db)):
    u=_require(request,db,mid,{SUPERVISOR_ROLE});_,s=access_student(db,request,mid,sid,True)
    row=db.get(AbsenceExcuse,eid)
    if not row or row.student_id!=sid:raise HTTPException(404,'العذر غير موجود')
    if row.status!='pending':raise HTTPException(409,'العذر تمت مراجعته')
    progress=db.scalar(select(HalaqaProgress).where(HalaqaProgress.student_id==sid,HalaqaProgress.day==row.day))
    if data.status=='approved' and progress and progress.attendance=='present':raise HTTPException(409,'الطالب مسجل حاضرًا؛ صحح سجل الحضور قبل قبول العذر')
    if data.status=='approved' and progress:
        audit_change(db,u.id,mid,'attendance_excused',sid,{'day':str(row.day),'attendance':progress.attendance},{'day':str(row.day),'attendance':'excused'})
        progress.attendance='excused'
    row.status=data.status;row.review_note=data.note;row.reviewed_by=u.id;row.updated_at=datetime.utcnow()
    audit(db,u.id,mid,'excuse_reviewed',f'{sid}:{eid}:{data.status}');db.commit();return {'ok':True}

class MessageInput(Input):
    body: str=Field(min_length=1,max_length=3000)

@router.post('/mosques/{mid}/students/{sid}/messages')
def send_message(mid:int,sid:int,data:MessageInput,request:Request,db=Depends(get_db)):
    u,s=access_student(db,request,mid,sid)
    if not s.active:raise HTTPException(409,'الطالب مؤرشف؛ المراسلات للقراءة فقط')
    db.add(StudentMessage(student_id=sid,sender_id=u.id,body=data.body));audit(db,u.id,mid,'student_message_sent',str(sid));db.commit();return {'ok':True}

@router.get('/mosques/{mid}/students/{sid}/messages')
def list_messages(mid:int,sid:int,request:Request,before:int=0,db=Depends(get_db)):
    u,s=access_student(db,request,mid,sid)
    q=select(StudentMessage).where(StudentMessage.student_id==sid)
    if before:q=q.where(StudentMessage.id<before)
    rows=db.scalars(q.order_by(StudentMessage.id.desc()).limit(100)).all()
    return {'messages':[{'id':m.id,'sender':db.get(User,m.sender_id).full_name,'mine':m.sender_id==u.id,'body':m.body,'created_at':m.created_at.isoformat()+'Z'} for m in reversed(rows)],'before':rows[-1].id if len(rows)==100 else None}

class ThreadReadInput(Input):
    last_id:int=Field(ge=0)

@router.post('/mosques/{mid}/students/{sid}/messages/read')
def read_thread(mid:int,sid:int,data:ThreadReadInput,request:Request,db=Depends(get_db)):
    u,_=access_student(db,request,mid,sid)
    last=db.get(StudentMessage,data.last_id)
    if not last or last.student_id!=sid:raise HTTPException(404,'الرسالة غير موجودة')
    row=db.scalar(select(ThreadRead).where(ThreadRead.student_id==sid,ThreadRead.user_id==u.id))
    if not row:row=ThreadRead(student_id=sid,user_id=u.id,last_id=0);db.add(row)
    row.last_id=max(row.last_id,data.last_id);db.commit();return {'ok':True}

class AnnouncementInput(Input):
    halaqa_id:int|None=None
    title:str=Field(min_length=2,max_length=180)
    body:str=Field(min_length=1,max_length=3000)
    expires:date|None=None

@router.post('/mosques/{mid}/announcements')
def announce(mid:int,data:AnnouncementInput,request:Request,db=Depends(get_db)):
    u=_require(request,db,mid,STAFF)
    if data.halaqa_id:ring_access(db,request,mid,data.halaqa_id)
    elif u.role!='owner':raise HTTPException(403,'الإعلان لجميع المسجد من صلاحية المالك')
    if data.expires and data.expires<today():raise HTTPException(422,'تاريخ الانتهاء في الماضي')
    row=Announcement(mosque_id=mid,created_by=u.id,**data.model_dump());db.add(row);audit(db,u.id,mid,'announcement_created',data.title);db.commit();return {'ok':True}

@router.post('/mosques/{mid}/announcements/{aid}/archive')
def archive_announcement(mid:int,aid:int,request:Request,db=Depends(get_db)):
    u=_require(request,db,mid,STAFF);row=db.get(Announcement,aid)
    if not row or row.mosque_id!=mid:raise HTTPException(404,'الإعلان غير موجود')
    if u.role!='owner':
        if row.halaqa_id not in _ring_ids_for(u,mid,db) or (u.role==TEACHER_ROLE and row.created_by!=u.id):raise HTTPException(403,'لا يمكنك أرشفة هذا الإعلان')
    row.active=False;audit(db,u.id,mid,'announcement_archived',str(aid));db.commit();return {'ok':True}

class Criterion(Input):
    name:str=Field(min_length=2,max_length=100)
    maximum:float=Field(gt=0,le=1000,allow_inf_nan=False)

class OrganizedInput(Input):
    title:str=Field(min_length=2,max_length=180)
    syllabus:str=Field(min_length=2,max_length=2000)
    due:date
    repeat_days:int=Field(0,ge=0,le=365)
    pass_percent:float=Field(70,ge=0,le=100,allow_inf_nan=False)
    criteria:list[Criterion]=Field(min_length=1,max_length=10)
    @model_validator(mode='after')
    def check(self):
        if len({c.name for c in self.criteria})!=len(self.criteria):raise ValueError('أسماء المعايير مكررة')
        if sum(c.maximum for c in self.criteria)>1000:raise ValueError('مجموع الدرجات يتجاوز 1000')
        return self

@router.post('/mosques/{mid}/students/{sid}/organized-exams')
def create_exam(mid:int,sid:int,data:OrganizedInput,request:Request,db=Depends(get_db)):
    u,_=access_student(db,request,mid,sid,True)
    if data.due<today():raise HTTPException(422,'موعد الاختبار في الماضي')
    d=data.model_dump();criteria=d.pop('criteria');row=OrganizedExam(student_id=sid,created_by=u.id,criteria_json=json.dumps(criteria,ensure_ascii=False),**d)
    db.add(row);db.flush();audit(db,u.id,mid,'organized_exam_created',str(row.id));db.commit();return {'ok':True,'id':row.id}

class ResultInput(Input):
    day:date
    scores:list[float]=Field(min_length=1,max_length=10)
    notes:str=Field('',max_length=2000)
    retake_due:date|None=None
    @model_validator(mode='after')
    def dates(self):
        import math
        if self.day>today():raise ValueError('النتيجة لا تسجل ليوم مستقبلي')
        if any(not math.isfinite(x) or x<0 for x in self.scores):raise ValueError('الدرجات غير صالحة')
        if self.retake_due and self.retake_due<=today():raise ValueError('موعد الإعادة يجب أن يكون مستقبلًا')
        return self

@router.put('/mosques/{mid}/students/{sid}/organized-exams/{eid}/result')
def assess_exam(mid:int,sid:int,eid:int,data:ResultInput,request:Request,db=Depends(get_db)):
    u,_=access_student(db,request,mid,sid,True);exam=db.get(OrganizedExam,eid)
    if not exam or exam.student_id!=sid:raise HTTPException(404,'الاختبار غير موجود')
    if exam.status=='cancelled':raise HTTPException(409,'الاختبار ملغى')
    if data.day<exam.due:raise HTTPException(422,'لا يمكن تسجيل نتيجة قبل موعد الاختبار')
    from app.models import ExamBooking
    booking=db.get(ExamBooking,eid)
    if booking and booking.approved:raise HTTPException(409,'اعتمد انتقال المستوى؛ لا يمكن تغيير النتيجة بعد الاعتماد')
    if booking and u.role==TEACHER_ROLE and u.id!=booking.examiner_id:raise HTTPException(403,'تسجيل النتيجة للمختبر المحدد أو المشرف')
    criteria=json.loads(exam.criteria_json)
    if len(data.scores)!=len(criteria) or any(v>c['maximum'] for v,c in zip(data.scores,criteria)):raise HTTPException(422,'الدرجات لا تطابق معايير الاختبار')
    total=sum(c['maximum'] for c in criteria);score=sum(data.scores);passed=100*score/total>=exam.pass_percent
    if not passed and not data.retake_due and not exam.next_exam_id:raise HTTPException(422,'حدد موعد الإعادة عند عدم الاجتياز')
    if passed and data.retake_due:raise HTTPException(422,'الاختبار مجتاز؛ اترك موعد الإعادة فارغًا')
    result=db.get(OrganizedResult,eid);before={'scores':json.loads(result.scores_json),'passed':result.passed} if result else {}
    legacy=db.get(StudentExam,result.legacy_exam_id) if result else StudentExam(student_id=sid,created_by=u.id)
    if not result:db.add(legacy)
    legacy.day=data.day;legacy.title=exam.title;legacy.score=score;legacy.total=total;legacy.notes=data.notes
    db.flush()
    if not result:result=OrganizedResult(exam_id=eid,legacy_exam_id=legacy.id,assessed_by=u.id);db.add(result)
    result.day=data.day;result.scores_json=json.dumps(data.scores);result.notes=data.notes;result.passed=passed;result.assessed_by=u.id
    # One follow-up per exam. Never silently overwrite a follow-up already assessed.
    next_due=data.retake_due if not passed else (data.day+timedelta(days=exam.repeat_days) if exam.repeat_days else None)
    following=db.get(OrganizedExam,exam.next_exam_id) if exam.next_exam_id else None
    if following and following.status=='scheduled':
        if next_due:following.due=next_due;following.title=exam.title
        else:following.status='cancelled'
    elif next_due and not following:
        following=OrganizedExam(student_id=sid,title=exam.title,syllabus=exam.syllabus,due=next_due,repeat_days=exam.repeat_days,pass_percent=exam.pass_percent,criteria_json=exam.criteria_json,created_by=u.id)
        db.add(following);db.flush();exam.next_exam_id=following.id
    exam.status='completed';audit_change(db,u.id,mid,'organized_exam_assessed',eid,before,{'scores':data.scores,'passed':passed,'day':str(data.day),'next_due':str(next_due) if next_due else None});db.commit()
    return {'ok':True,'passed':passed}

@router.post('/mosques/{mid}/students/{sid}/organized-exams/{eid}/cancel')
def cancel_exam(mid:int,sid:int,eid:int,request:Request,db=Depends(get_db)):
    u,_=access_student(db,request,mid,sid,True);row=db.get(OrganizedExam,eid)
    if not row or row.student_id!=sid:raise HTTPException(404,'الاختبار غير موجود')
    if row.status!='scheduled':raise HTTPException(409,'يمكن إلغاء الاختبار القادم فقط')
    row.status='cancelled';audit(db,u.id,mid,'organized_exam_cancelled',str(eid));db.commit();return {'ok':True}

class ArchiveInput(Input):
    active:bool
    reason:str=Field(min_length=2,max_length=500)

@router.put('/mosques/{mid}/students/{sid}/archive')
def archive_student(mid:int,sid:int,data:ArchiveInput,request:Request,db=Depends(get_db)):
    u=_require(request,db,mid,{SUPERVISOR_ROLE});row=db.get(HalaqaStudent,sid)
    if not row or row.halaqa_id not in _ring_ids_for(u,mid,db):raise HTTPException(404,'الطالب غير موجود')
    before=row.active;row.active=data.active
    audit_change(db,u.id,mid,'student_archive_changed',sid,{'active':before},data.model_dump());db.commit();return {'ok':True}

@router.put('/mosques/{mid}/users/{uid}/archive')
def archive_user(mid:int,uid:int,data:ArchiveInput,request:Request,db=Depends(get_db)):
    u=_require(request,db,mid,{'owner'});row=db.get(User,uid)
    if not row or row.mosque_id!=mid or row.role=='owner':raise HTTPException(404,'الحساب غير قابل للأرشفة')
    before=row.active;row.active=data.active
    if not data.active:db.execute(delete(AuthSession).where(AuthSession.user_id==uid))
    audit_change(db,u.id,mid,'account_archive_changed',uid,{'active':before},data.model_dump());db.commit();return {'ok':True}

@router.get('/mosques/{mid}/audit')
def audit_history(mid:int,request:Request,before:int=0,db=Depends(get_db)):
    _require(request,db,mid,{'owner'});q=select(AuditLog).where(AuditLog.mosque_id==mid)
    if before:q=q.where(AuditLog.id<before)
    rows=db.scalars(q.order_by(AuditLog.id.desc()).limit(100)).all()
    return {'rows':[{'id':r.id,'actor':db.get(User,r.user_id).full_name if r.user_id else 'النظام','action':r.action,'details':r.details,'at':r.created_at.isoformat()+'Z'} for r in rows],'before':rows[-1].id if len(rows)==100 else None}

@router.get('/mosques/{mid}/students/{sid}/development')
def student_development(mid:int,sid:int,request:Request,db=Depends(get_db)):
    u,s=access_student(db,request,mid,sid)
    records={r.surah:r for r in db.scalars(select(QuranRecord).where(QuranRecord.student_id==sid))}
    quran=[{'number':i,'name':name,'status':records[i].status if i in records else 'unstarted','notes':records[i].notes if i in records else '', 'approved_by':db.get(User,records[i].approved_by).full_name if i in records else '', 'updated_at':records[i].updated_at.isoformat()+'Z' if i in records else ''} for i,name in enumerate(SURAH_NAMES,1)]
    excuses=db.scalars(select(AbsenceExcuse).where(AbsenceExcuse.student_id==sid).order_by(AbsenceExcuse.day.desc())).all()
    exams=[]
    for x in db.scalars(select(OrganizedExam).where(OrganizedExam.student_id==sid).order_by(OrganizedExam.due.desc(),OrganizedExam.id.desc())):
        r=db.get(OrganizedResult,x.id)
        exams.append({'id':x.id,'title':x.title,'syllabus':x.syllabus,'due':str(x.due),'repeat_days':x.repeat_days,'criteria':json.loads(x.criteria_json),'pass_percent':x.pass_percent,'status':x.status,'next_exam_id':x.next_exam_id,'result':{'day':str(r.day),'scores':json.loads(r.scores_json),'notes':r.notes,'passed':r.passed} if r else None})
    return {'quran':quran,'quran_states':QURAN_STATES,'excuses':[{'id':x.id,'day':str(x.day),'reason':x.reason,'status':x.status,'review_note':x.review_note} for x in excuses], 'exams':exams}

@router.get('/mosques/{mid}/development')
def development_dashboard(mid:int,request:Request,db=Depends(get_db)):
    u=_require(request,db,mid,PORTAL_ROLES);ids=_ring_ids_for(u,mid,db);students=active_students(db,u,mid);sids=[s.id for s in students]
    calendars=[]
    for rid in ids:
        r=db.get(RingCalendar,rid);ring=db.get(Halaqa,rid)
        exceptions=db.scalars(select(CalendarException).where(CalendarException.halaqa_id==rid,CalendarException.day>=today()-timedelta(days=30)).order_by(CalendarException.day).limit(200)).all()
        calendars.append({'id':rid,'name':ring.name,'configured':r is not None,'weekdays':json.loads(r.weekdays_json) if r else [],'time_text':r.time_text if r else '', 'today_held':meeting_day(db,rid,today()),'exceptions':[{'day':str(x.day),'held':x.held,'reason':x.reason} for x in exceptions]})
    anns=db.scalars(select(Announcement).where(Announcement.mosque_id==mid,Announcement.active.is_(True),or_(Announcement.halaqa_id.is_(None),Announcement.halaqa_id.in_(ids)),or_(Announcement.expires.is_(None),Announcement.expires>=today())).order_by(Announcement.id.desc()).limit(200)).all()
    unread=[];pending=[]
    for s in students:
        r=db.scalar(select(ThreadRead).where(ThreadRead.student_id==s.id,ThreadRead.user_id==u.id))
        count=db.scalar(select(func.count(StudentMessage.id)).where(StudentMessage.student_id==s.id,StudentMessage.id>(r.last_id if r else 0),StudentMessage.sender_id!=u.id))
        if count:unread.append({'id':s.id,'name':s.full_name,'count':count})
    if u.role in ('owner',SUPERVISOR_ROLE):
        pending=[{'id':x.id,'student_id':x.student_id,'name':db.get(HalaqaStudent,x.student_id).full_name,'day':str(x.day),'reason':x.reason} for x in db.scalars(select(AbsenceExcuse).where(AbsenceExcuse.student_id.in_(sids),AbsenceExcuse.status=='pending'))]
    archives=[{'id':s.id,'name':s.full_name,'halaqa_id':s.halaqa_id} for s in db.scalars(select(HalaqaStudent).where(HalaqaStudent.halaqa_id.in_(ids),HalaqaStudent.active.is_(False)))] if u.role in ('owner',SUPERVISOR_ROLE) else []
    upcoming=[{'id':e.id,'student_id':e.student_id,'name':db.get(HalaqaStudent,e.student_id).full_name,'title':e.title,'due':str(e.due)} for e in db.scalars(select(OrganizedExam).where(OrganizedExam.student_id.in_(sids),OrganizedExam.status=='scheduled').order_by(OrganizedExam.due))]
    return {'calendars':calendars,'announcements':[{'id':a.id,'title':a.title,'body':a.body,'ring':db.get(Halaqa,a.halaqa_id).name if a.halaqa_id else 'جميع حلقات المسجد','expires':str(a.expires) if a.expires else '', 'can_archive':u.role=='owner' or (a.halaqa_id in ids and (u.role==SUPERVISOR_ROLE or a.created_by==u.id))} for a in anns], 'unread':unread,'pending_excuses':pending,'archived_students':archives,'upcoming_exams':upcoming}

# Excel import stages validated requests; only the existing owner approval creates accounts.
IMPORT_HEADERS=['اسم الطالب','المتابع','اسم ولي الأمر','جوال ولي الأمر','صلة القرابة','تاريخ الميلاد','الصف الدراسي','الحفظ الحالي','ملاحظات']
IMPORT_FIELDS=['full_name','recipient_type','guardian_name','guardian_phone','guardian_relation','birth_date','school_grade','current_memorization','notes']

def normalized_name(value):
    import unicodedata
    text=unicodedata.normalize('NFKC',value or '')
    return ' '.join(''.join(c for c in text if not unicodedata.combining(c) and c!='ـ').split()).casefold()

def known_names(db,mid):
    names=list(db.scalars(select(HalaqaStudent.full_name).join(Halaqa,Halaqa.id==HalaqaStudent.halaqa_id).where(Halaqa.mosque_id==mid)))
    names+=list(db.scalars(select(HalaqaAccountRequest.full_name).where(HalaqaAccountRequest.mosque_id==mid,HalaqaAccountRequest.request_type=='student',HalaqaAccountRequest.status.in_(['pending','returned']))))
    return {normalized_name(n) for n in names}

@router.get('/mosques/{mid}/imports/template.xlsx')
def excel_template(mid:int,request:Request,db=Depends(get_db)):
    _require(request,db,mid,STAFF)
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    from openpyxl.worksheet.datavalidation import DataValidation
    from openpyxl.utils import get_column_letter
    w=Workbook();s=w.active;s.title='الطلاب';s.sheet_view.rightToLeft=True;s.append(IMPORT_HEADERS);s.freeze_panes='A2'
    for c in s[1]:c.font=Font(bold=True,color='FFFFFF');c.fill=PatternFill('solid',fgColor='145344')
    for i in range(1,10):s.column_dimensions[get_column_letter(i)].width=24
    for row in s.iter_rows(min_row=2,max_row=301,min_col=1,max_col=9):
        for c in row:c.number_format='@'
    dv=DataValidation(type='list',formula1='"الطالب,ولي الأمر"');s.add_data_validation(dv);dv.add('B2:B301')
    guide=w.create_sheet('تعليمات');guide.sheet_view.rightToLeft=True;guide.column_dimensions['A'].width=95
    for text in ['املأ ورقة الطلاب. الحد الأقصى 300 طالب لكل ملف.','المتابع: الطالب أو ولي الأمر. اسم ولي الأمر مطلوب عند اختياره.','تاريخ الميلاد اختياري بصيغة YYYY-MM-DD؛ احفظ الجوال كنص للحفاظ على الصفر.','الاستيراد ينشئ طلبات للمالك بعد المعاينة ولا يصدر حسابات تلقائيًا.','التطابق في اسم الطالب داخل المسجد ينبه لاحتمال التكرار ويوقف ذلك الصف؛ سجله يدويًا بعد التحقق إن كان شخصًا مختلفًا.']:guide.append([text])
    out=io.BytesIO();w.save(out);return Response(out.getvalue(),media_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',headers={'Content-Disposition':'attachment; filename="students-template.xlsx"'})

@router.post('/mosques/{mid}/imports/preview')
async def excel_preview(mid:int,request:Request,halaqa_id:int=Form(...),file:UploadFile=File(...),db=Depends(get_db)):
    u=ring_access(db,request,mid,halaqa_id)
    raw=await file.read(2*1024*1024+1)
    if len(raw)>2*1024*1024 or not (file.filename or '').lower().endswith('.xlsx'):raise HTTPException(422,'ارفع ملف xlsx لا يتجاوز 2 ميجابايت')
    from openpyxl import load_workbook
    try:
        with zipfile.ZipFile(io.BytesIO(raw)) as z:
            if len(z.infolist())>100 or sum(i.file_size for i in z.infolist())>12*1024*1024:raise ValueError('large archive')
            if any('vbaproject' in i.filename.lower() for i in z.infolist()):raise ValueError('macros')
        w=load_workbook(io.BytesIO(raw),read_only=True,data_only=False,keep_links=False)
        sheet=w['الطلاب'] if 'الطلاب' in w.sheetnames else w.worksheets[0]
        rows=list(sheet.iter_rows(max_row=302,max_col=10));w.close()
    except Exception:raise HTTPException(422,'ملف Excel غير صالح أو يتجاوز الحد؛ استخدم القالب')
    if not rows or [c.value for c in rows[0][:9]]!=IMPORT_HEADERS:raise HTTPException(422,'عناوين الأعمدة لا تطابق القالب')
    if any(c.value is not None for c in rows[0][9:]):raise HTTPException(422,'أعمدة زائدة؛ استخدم القالب')
    existing=known_names(db,mid);seen=set();preview=[]
    for line,cells in enumerate(rows[1:],2):
        if not any(c.value is not None for c in cells):continue
        if len(preview)>=300 or line>301:raise HTTPException(422,'الحد الأقصى 300 صف طالب')
        payload={'halaqa_id':halaqa_id};errors=[]
        for key,c in zip(IMPORT_FIELDS,cells[:9]):
            val=c.value
            if c.data_type=='f':errors.append('الصيغ غير مسموحة؛ استخدم قيمًا ثابتة')
            if key=='birth_date':
                if isinstance(val,datetime):val=val.date().isoformat()
                elif isinstance(val,date):val=val.isoformat()
                else:val=str(val).strip() if val else None
            else:val=str(val).strip() if val is not None else ''
            payload[key]=val
        if any(c.value is not None for c in cells[9:]):errors.append('عمود زائد')
        payload['recipient_type']={'الطالب':'student','الطالب نفسه':'student','ولي الأمر':'guardian','ولي الامر':'guardian'}.get(payload['recipient_type'],payload['recipient_type'])
        key=normalized_name(payload['full_name'])
        if key in seen or key in existing:errors.append('اسم مكرر في الملف أو في سجلات المسجد؛ راجع الطالب')
        seen.add(key)
        try:
            valid=StudentRequestInput(**payload)
            if valid.recipient_type=='guardian' and not valid.guardian_name:errors.append('اسم ولي الأمر مطلوب')
            payload=valid.model_dump(mode='json')
        except ValidationError as ex:errors+=['تحقق من '+str(e['loc'][0])+': '+e['msg'] for e in ex.errors()]
        preview.append({'row':line,'data':payload,'errors':errors})
    if not preview:raise HTTPException(422,'لا يوجد طلاب في الملف')
    token=secrets.token_urlsafe(24)
    db.execute(delete(ImportPreview).where(ImportPreview.created_at<datetime.utcnow()-timedelta(hours=2)))
    db.add(ImportPreview(token=token,mosque_id=mid,user_id=u.id,halaqa_id=halaqa_id,rows_json=json.dumps(preview,ensure_ascii=False)))
    db.commit();return {'token':token,'rows':preview}

class ImportCommitInput(Input):
    token:str=Field(min_length=20,max_length=64)
    rows:list[int]=Field(min_length=1,max_length=300)

@router.post('/mosques/{mid}/imports/commit')
def excel_commit(mid:int,data:ImportCommitInput,request:Request,db=Depends(get_db)):
    u=_require(request,db,mid,STAFF);batch=db.get(ImportPreview,data.token)
    if not batch or batch.mosque_id!=mid or batch.user_id!=u.id:raise HTTPException(404,'المعاينة غير موجودة')
    ring_access(db,request,mid,batch.halaqa_id)
    if batch.consumed or batch.created_at<datetime.utcnow()-timedelta(hours=2):raise HTTPException(409,'المعاينة مستخدمة أو منتهية؛ أعد رفع الملف')
    if len(set(data.rows))!=len(data.rows):raise HTTPException(422,'صفوف مكررة')
    source={r['row']:r for r in json.loads(batch.rows_json)};existing=known_names(db,mid);selected=[]
    for number in data.rows:
        row=source.get(number)
        if not row or row['errors']:raise HTTPException(422,'حدد صفوفًا سليمة من المعاينة فقط')
        obj=StudentRequestInput(**row['data']);key=normalized_name(obj.full_name)
        if key in existing:raise HTTPException(409,'ظهر اسم مكرر منذ المعاينة؛ أعد رفع الملف')
        existing.add(key);selected.append(obj)
    for obj in selected:
        db.add(HalaqaAccountRequest(mosque_id=mid,request_type='student',requested_by=u.id,status='pending',**obj.model_dump()))
    batch.consumed=True;audit(db,u.id,mid,'students_imported',str(len(selected)));db.commit()
    return {'ok':True,'count':len(selected)}
