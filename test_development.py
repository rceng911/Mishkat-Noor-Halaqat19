import io,json
from datetime import timedelta
from openpyxl import Workbook,load_workbook
from sqlalchemy import select
from app.api.student_file import today
from app.db.session import SessionLocal
from app.models import User,StudentExam,OrganizedExam,HalaqaStudent
from app.api.development import IMPORT_HEADERS
BASE='/api/halaqat/mosques/1'

def test_quran_map_scope_and_audit(network):
    n=network;sid=n['ids'][0];url=BASE+f'/students/{sid}'
    result=n['students'][0].get(url+'/development');assert result.status_code==200,result.text
    catalog=result.json()['quran'];assert len(catalog)==114
    assert catalog[2]['name']=='آل عمران' and catalog[-1]['name']=='الناس'
    assert n['teachers'][1].put(url+'/quran/1',json={'status':'memorized'}).status_code==404
    assert n['students'][0].put(url+'/quran/1',json={'status':'memorized'}).status_code==403
    assert n['teachers'][0].put(url+'/quran/115',json={'status':'memorized'}).status_code==422
    assert n['teachers'][0].put(url+'/quran/1',json={'status':'memorized','notes':'تسميع كامل'}).status_code==200
    assert n['students'][0].get(url+'/development').json()['quran'][0]['status']=='memorized'
    assert n['students'][1].get(url+'/development').status_code==404
    assert n['teachers'][0].get(BASE+'/audit').status_code==403
    rows=n['owner'].get(BASE+'/audit').json()['rows']
    changed=next(r for r in rows if r['action']=='quran_updated')
    assert json.loads(changed['details'])['before']['status']=='unstarted'
    assert json.loads(changed['details'])['after']['status']=='memorized'


def test_calendar_actual_meetings_and_excuse_flow(network):
    n=network;t=n['teachers'][0];s=n['students'][0];sup=n['supervisors'][0];rid=n['rings'][0];sid=n['ids'][0];url=BASE+f'/students/{sid}';cal=BASE+f'/rings/{rid}'
    assert sup.get(BASE+'/monitor').json()['missing_attendance']==[]
    assert t.put(cal+'/calendar',json={'weekdays':[today().weekday()]}).status_code==403
    assert sup.put(cal+'/calendar',json={'weekdays':[today().weekday()],'time_text':'بعد العصر'}).status_code==200
    assert len(sup.get(BASE+'/monitor').json()['missing_attendance'])==1
    assert sup.put(cal+'/calendar-exception',json={'day':str(today()),'held':False,'reason':'إجازة'}).status_code==200
    assert sup.get(BASE+'/monitor').json()['missing_attendance']==[]
    assert s.post(url+'/excuses',json={'day':str(today()),'reason':'ظرف عائلي'}).status_code==422
    assert sup.delete(cal+'/calendar-exception/'+str(today())).status_code==200
    assert t.post(url+'/progress',json={'day':str(today()),'attendance':'absent'}).status_code==200
    assert s.post(url+'/excuses',json={'day':str(today()),'reason':'ظرف عائلي'}).status_code==200
    eid=s.get(url+'/development').json()['excuses'][0]['id']
    assert t.put(url+f'/excuses/{eid}',json={'status':'approved'}).status_code==403
    assert n['supervisors'][1].get(url+'/file').status_code==200
    assert sup.put(url+f'/excuses/{eid}',json={'status':'approved','note':'مقبول'}).status_code==200
    assert s.get(url+'/file').json()['progress'][0]['attendance']=='excused'
    # Later bulk absence must respect the accepted excuse.
    assert t.post(BASE+'/attendance/batch',json={'day':str(today()),'rows':[{'student_id':sid,'attendance':'absent'}]}).status_code==200
    assert s.get(url+'/file').json()['progress'][0]['attendance']=='excused'
    assert s.post(url+'/excuses',json={'day':str(today()),'reason':'عذر مكرر'}).status_code==409


def test_messages_announcements_unread_and_cross_family(network):
    n=network;t=n['teachers'][0];s=n['students'][0];url=BASE+f'/students/{n["ids"][0]}'
    assert s.post(url+'/messages',json={'body':'أريد متابعة مراجعة ابني'}).status_code==200
    assert n['students'][1].get(url+'/messages').status_code==404
    assert n['teachers'][1].post(url+'/messages',json={'body':'غير مصرح'}).status_code==404
    msg=t.get(url+'/messages').json()['messages'][0]
    assert t.get(BASE+'/development').json()['unread'][0]['count']==1
    assert t.post(url+'/messages/read',json={'last_id':msg['id']}).status_code==200
    assert t.get(BASE+'/development').json()['unread']==[]
    assert t.post(url+'/messages',json={'body':'نراجع معه غدًا'}).status_code==200
    assert s.get(BASE+'/development').json()['unread'][0]['count']==1
    assert t.post(BASE+'/announcements',json={'title':'إعلان عام','body':'اختبار'}).status_code==403
    assert t.post(BASE+'/announcements',json={'halaqa_id':n['rings'][1],'title':'إعلان حلقة','body':'اختبار'}).status_code==404
    assert t.post(BASE+'/announcements',json={'halaqa_id':n['rings'][0],'title':'إعلان حلقة','body':'اختبار'}).status_code==200
    assert len(s.get(BASE+'/development').json()['announcements'])==1
    assert n['students'][1].get(BASE+'/development').json()['announcements']==[]
    assert n['owner'].post(BASE+'/announcements',json={'title':'لكل الحلقات','body':'مرحبًا'}).status_code==200
    assert len(n['students'][1].get(BASE+'/development').json()['announcements'])==1


def test_organized_exams_results_retake_repeat_and_no_duplicates(network):
    n=network;t=n['teachers'][0];s=n['students'][0];url=BASE+f'/students/{n["ids"][0]}'
    payload={'title':'الاختبار الدوري','syllabus':'جزء عم','due':str(today()),'repeat_days':30,'criteria':[{'name':'الحفظ','maximum':60},{'name':'التجويد','maximum':40}]}
    assert s.post(url+'/organized-exams',json=payload).status_code==403
    eid=t.post(url+'/organized-exams',json=payload).json()['id']
    result={'day':str(today()),'scores':[20,10],'notes':'يحتاج تثبيتًا'}
    assert t.put(url+f'/organized-exams/{eid}/result',json=result).status_code==422
    result['retake_due']=str(today()+timedelta(days=7))
    assert t.put(url+f'/organized-exams/{eid}/result',json=result).status_code==200
    assert t.put(url+f'/organized-exams/{eid}/result',json=result).status_code==200
    exams=s.get(url+'/development').json()['exams'];assert len(exams)==2
    original=next(e for e in exams if e['id']==eid);assert not original['result']['passed']
    assert len(s.get(url+'/file').json()['exams'])==1
    assert t.put(url+f'/organized-exams/{eid}/result',json=result|{'scores':[60,38],'retake_due':None}).status_code==200
    exams=s.get(url+'/development').json()['exams'];assert len(exams)==2
    upcoming=next(e for e in exams if e['status']=='scheduled');assert upcoming['due']==str(today()+timedelta(days=30))
    assert len(s.get(url+'/file').json()['exams'])==1
    assert n['teachers'][1].post(url+f'/organized-exams/{upcoming["id"]}/cancel').status_code==404
    assert t.post(url+f'/organized-exams/{upcoming["id"]}/cancel').status_code==200


def test_archiving_preserves_students_and_sessions(network):
    n=network;o=n['owner'];s=n['students'][0];sid=n['ids'][0];url=BASE+f'/students/{sid}'
    assert n['teachers'][0].post(url+'/progress',json={'day':str(today()),'memorized':'الفاتحة','memorized_amount':1}).status_code==200
    assert n['teachers'][0].put(url+'/archive',json={'active':False,'reason':'انتقل'}).status_code==403
    assert n['supervisors'][0].put(url+'/archive',json={'active':False,'reason':'انتقل'}).status_code==200
    assert s.get(url+'/file').status_code==404
    old=o.get(url+'/file').json();assert not old['can_edit'] and old['progress'][0]['memorized']=='الفاتحة'
    assert o.post(url+'/assignments',json={'due':str(today()),'memorization':'الناس'}).status_code==404
    assert len(o.get(BASE+'/development').json()['archived_students'])==1
    assert o.put(url+'/archive',json={'active':True,'reason':'عودة'}).status_code==200
    assert s.get(url+'/file').json()['progress'][0]['memorized']=='الفاتحة'
    uid=s.get('/api/halaqat/context').json()['user']['id']
    assert o.put(BASE+f'/users/{uid}/archive',json={'active':False,'reason':'إيقاف دخول'}).status_code==200
    assert s.get('/api/halaqat/context').status_code==401
    assert o.put(BASE+f'/users/{uid}/archive',json={'active':True,'reason':'عودة'}).status_code==200
    assert s.get('/api/halaqat/context').status_code==401 # old revoked token stays revoked
    own=o.get('/api/halaqat/context').json()['user']['id']
    assert o.put(BASE+f'/users/{own}/archive',json={'active':False,'reason':'اختبار'}).status_code==404


def workbook(rows):
    w=Workbook();s=w.active;s.title='الطلاب';s.append(IMPORT_HEADERS)
    for row in rows:s.append(row)
    b=io.BytesIO();w.save(b);return b.getvalue()


def test_excel_preview_errors_commit_review_and_replay(network):
    n=network;t=n['teachers'][0];o=n['owner'];rid=n['rings'][0]
    raw=workbook([['طالب للاستيراد','ولي الأمر','والد الطالب','0501234567','الأب','2014-05-02','الخامس','الفاتحة','ملاحظات'],['طالب للاستيراد','الطالب'],['طالب بلا ولي','ولي الأمر'],['=1+2','الطالب']])
    def preview(client,ring=rid):return client.post(BASE+'/imports/preview',data={'halaqa_id':ring},files={'file':('students.xlsx',raw,'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')})
    assert preview(n['students'][0]).status_code==403
    assert preview(t,n['rings'][1]).status_code==404
    r=preview(t);assert r.status_code==200,r.text;d=r.json()
    assert not d['rows'][0]['errors'];assert all(row['errors'] for row in d['rows'][1:])
    assert d['rows'][0]['data']['guardian_phone']=='0501234567'
    assert o.post(BASE+'/imports/commit',json={'token':d['token'],'rows':[2]}).status_code==404
    assert t.post(BASE+'/imports/commit',json={'token':d['token'],'rows':[2,3]}).status_code==422
    assert t.post(BASE+'/imports/commit',json={'token':d['token'],'rows':[2]}).status_code==200
    assert t.post(BASE+'/imports/commit',json={'token':d['token'],'rows':[2]}).status_code==409
    req=next(r for r in o.get(BASE+'/dashboard').json()['requests'] if r['full_name']=='طالب للاستيراد');assert req['status']=='pending'
    assert len(o.get(BASE+'/dashboard').json()['students'])==2
    assert o.post(BASE+f'/requests/{req["id"]}/approve',json={'username':'اب_المستورد'}).status_code==200
    assert len(o.get(BASE+'/dashboard').json()['students'])==3
    template=t.get(BASE+'/imports/template.xlsx');assert template.status_code==200
    assert load_workbook(io.BytesIO(template.content)).worksheets[0]['A1'].value=='اسم الطالب'


def test_multi_mosque_isolation_and_owner_views(network):
    n=network;o=n['owner'];mid=o.post('/api/halaqat/organizations',json={'name':'مسجد آخر'}).json()['id'];other=f'/api/halaqat/mosques/{mid}'
    assert n['supervisors'][0].post('/api/halaqat/organizations',json={'name':'غير مصرح'}).status_code==403
    assert n['students'][0].get(other+'/dashboard').status_code==403
    assert n['teachers'][0].get(other+'/development').status_code==403
    assert o.get(other+'/development',headers={'X-Portal-View':'supervisor'}).status_code==200
    assert o.get(other+'/dashboard',headers={'X-Portal-View':'student'}).json()['students']==[]
    sup=o.post(other+'/supervisors',json={'full_name':'مشرف المسجد الجديد','username':'مشرف_الجديد'}).json()['user']['id']
    rid=o.post(other+'/rings',json={'name':'حلقة المسجد الجديد','supervisor_ids':[sup]}).json()['ring']['id']
    assert n['teachers'][0].post(BASE+'/requests/student',json={'full_name':'تجاوز','halaqa_id':rid,'recipient_type':'student'}).status_code==400
    assert len(o.get('/api/halaqat/organizations/overview').json()['mosques'])==2
    assert len(n['students'][0].get('/api/halaqat/context').json()['mosques'])==1
    own=o.get('/api/halaqat/context').json()['user']['id']
    req=o.post(other+'/requests/student',json={'full_name':'محمد في مسجد آخر','halaqa_id':rid,'recipient_type':'student'}).json()['request']['id']
    assert o.post(other+f'/requests/{req}/approve',json={'existing_user_id':own}).status_code==200
    assert len(o.get(other+'/dashboard',headers={'X-Portal-View':'student'}).json()['students'])==1
    assert o.get(other+'/audit').json()['rows']
    assert n['supervisors'][0].get(other+'/audit').status_code==403
